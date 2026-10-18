"""Bulk product import from one Excel workbook (pictures are embedded in it).

Two passes over the same code: `run(..., dry_run=True)` only validates and reports; `run(..., dry_run=False)` validates and,
when there is not a single error, writes everything in ONE transaction (all or nothing). Related records (brand, category,
warehouse, price list) are matched by NAME and are never created here, so a typo is an error, not a new record.

The layout is defined in `import_template.py`. Pictures come from the workbook itself (put in the "Images" column, see
`import_images.py`); there is no separate upload.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.core.files.base import ContentFile
from django.db import transaction
from openpyxl import load_workbook
from rest_framework import status

from apps.core import audit
from apps.core.exceptions import ServiceError

from . import services
from .import_template import (
    COMPARE_PREFIX, NO, PATH_SEPARATOR, PRICE_PREFIX, PRODUCT_COLUMNS, SHEET_PRODUCTS, YES,
)
from .import_images import floating_pictures, in_cell_pictures
from .models import Brand, Category, PriceList, ProductVariant

MAX_ROWS = 2000
MAX_XLSX_BYTES = 60 * 1024 * 1024  # pictures live inside the workbook, so it can be large
MAX_XLSX_UNPACKED = 400 * 1024 * 1024
MAX_XLSX_PARTS = 5000
MAX_IMAGES_PER_PRODUCT = 10
MAX_REPORTED = 200  # errors / warnings returned to the client (the totals are always exact)

_TRUE = {"نعم", "yes", "y", "true", "1", "صح"}
_FALSE = {"لا", "no", "n", "false", "0", "خطأ"}
_EXCEL_ERRORS = {"#VALUE!", "#N/A", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#DIV/0!", "#SPILL!", "#CALC!", "#GETTING_DATA"}
_ATTR_SPLIT = re.compile(r"[,\n؛;،]+")

PRODUCT_LEVEL = ["name", "name_ar", "name_en", "name_he", "description", "description_ar", "description_en",
                 "description_he", "brand", "category", "publish", "track"]


@dataclass
class Issue:
    code: str
    detail: str
    sheet: str = ""
    row: int | None = None
    column: str = ""
    value: str = ""

    def as_dict(self):
        return {"sheet": self.sheet, "row": self.row, "column": self.column, "code": self.code,
                "detail": self.detail, "value": self.value}


@dataclass
class Report:
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    products: int = 0
    variants: int = 0
    images: int = 0
    stock_lines: int = 0
    created: int = 0

    def error(self, code, detail, **kw):
        self.errors.append(Issue(code, detail, **kw))

    def warn(self, code, detail, **kw):
        self.warnings.append(Issue(code, detail, **kw))

    def as_dict(self, dry_run):
        return {
            "dry_run": dry_run,
            "valid": not self.errors,
            "summary": {"products": self.products, "variants": self.variants, "images": self.images,
                        "stock_lines": self.stock_lines},
            "errors": [e.as_dict() for e in self.errors[:MAX_REPORTED]], "errors_total": len(self.errors),
            "warnings": [w.as_dict() for w in self.warnings[:MAX_REPORTED]], "warnings_total": len(self.warnings),
        }


# ---------------------------------------------------------------- cell helpers


def _text(v):
    """A cell as clean text. Numbers typed into text columns (barcodes, SKUs) must not become '6.281E+12' or '12.0'."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return YES if v else NO
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    text = str(v).strip()
    # an Excel error value (#VALUE! ...) is what a cell holding a picture reads as; it is never real data
    return "" if text in _EXCEL_ERRORS else text


def _norm(s):
    return " ".join(str(s).split()).casefold()


def _label_key(label):
    return _norm(str(label).replace("*", ""))


@dataclass
class ImageRef:
    """A picture of a product, taken from the workbook (`label` says where, for messages)."""

    label: str
    data: bytes


@dataclass
class Row:
    n: int
    cells: dict


class Context:
    """Everything looked up once: related records by name, existing SKUs/barcodes."""

    def __init__(self):
        self.brands = {}
        for b in Brand.objects.all():
            for name in {b.name, b.name_ar, b.name_en, b.name_he}:
                if name:
                    self.brands.setdefault(_norm(name), []).append(b)
        self.categories = list(Category.objects.all())
        self.cat_by_pk = {c.pk: c for c in self.categories}
        from apps.inventory.models import Warehouse

        self.warehouses = {}
        for w in Warehouse.objects.all():
            for name in {w.name, w.code}:
                self.warehouses.setdefault(_norm(name), []).append(w)
        self.price_lists = {_norm(p.name): p for p in PriceList.objects.select_related("currency")}
        self.sku_taken = {s.casefold() for s in ProductVariant.objects.values_list("sku", flat=True)}
        self.barcode_taken = set(ProductVariant.objects.exclude(barcode="").values_list("barcode", flat=True))
        self.has_default_price_list = PriceList.objects.filter(is_default=True, is_active=True).exists()

    # -- related records
    def find_brand(self, text):
        found = {b.pk: b for b in self.brands.get(_norm(text), [])}
        return list(found.values())

    def find_category(self, text):
        parts = [p.strip() for p in text.split(PATH_SEPARATOR)]
        if not all(parts):
            return []
        want_last = _norm(parts[-1])
        matches = []
        for c in self.categories:
            if want_last not in {_norm(n) for n in (c.name, c.name_ar, c.name_en, c.name_he) if n}:
                continue
            node, ok = c, True
            for part in reversed(parts[:-1]):
                node = self.cat_by_pk.get(node.parent_id)
                if node is None or _norm(part) not in {_norm(n) for n in (node.name, node.name_ar, node.name_en, node.name_he) if n}:
                    ok = False
                    break
            if ok:
                matches.append(c)
        return matches

    def find_warehouse(self, text):
        found = {w.pk: w for w in self.warehouses.get(_norm(text), [])}
        return list(found.values())

    # -- pictures
    def check_bytes(self, data):
        """None when `data` is a usable picture, else (code, detail)."""
        if len(data) > settings.MAX_UPLOAD_IMAGE_BYTES:
            return ("image_too_large", f"The picture is larger than {settings.MAX_UPLOAD_IMAGE_BYTES // (1024 * 1024)} MB.")
        try:
            services.validate_image(ContentFile(data, name="picture"))
        except ServiceError as e:
            return (e.code, e.detail)
        except Exception:
            return ("invalid_image", "The picture could not be read.")
        return None


# ---------------------------------------------------------------- reading the workbook


def _open_workbook(file):
    """-> (workbook, raw bytes). Opened fully (not read-only): embedded pictures are only visible that way."""
    name = getattr(file, "name", "") or ""
    if not name.lower().endswith(".xlsx"):
        raise ServiceError("invalid_file", "The file must be an Excel .xlsx workbook (use the downloaded template).")
    if getattr(file, "size", 0) > MAX_XLSX_BYTES:
        raise ServiceError("file_too_large", f"The Excel file can be at most {MAX_XLSX_BYTES // (1024 * 1024)} MB.")
    raw = file.read()
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            infos = z.infolist()
            if len(infos) > MAX_XLSX_PARTS or sum(i.file_size for i in infos) > MAX_XLSX_UNPACKED:
                raise ServiceError("file_too_large", "The Excel file is too big once unpacked.")
        return load_workbook(io.BytesIO(raw), data_only=True), raw
    except ServiceError:
        raise
    except Exception:
        raise ServiceError("invalid_file", "The file could not be read as an Excel workbook.")


def _sheet(wb, *names):
    wanted = {_norm(n) for n in names}
    for title in wb.sheetnames:
        if _norm(title) in wanted:
            return wb[title]
    return None


def _read_rows(ws):
    """[(row number, [cells])] without fully empty rows; the first row is the header."""
    out = []
    for n, values in enumerate(ws.iter_rows(values_only=True), start=1):
        if n > MAX_ROWS + 1:
            raise ServiceError("too_many_rows", f"The sheet '{ws.title}' has more than {MAX_ROWS} rows.")
        if any(v not in (None, "") and str(v).strip() != "" for v in values):
            out.append((n, list(values)))
    return out


def _header_map(rows, columns, sheet, report, extra_prefixes=()):
    """{column index: key}; price columns map to ('price'|'compare', PriceList name). Reports missing / unknown headers."""
    if not rows or rows[0][0] != 1:
        report.error("missing_header", "The first row must contain the column headers (use the downloaded template).", sheet=sheet)
        return None
    by_label = {}
    for c in columns:
        by_label[_label_key(c.label)] = c.key
        by_label[_norm(c.key)] = c.key
    mapping, seen = {}, set()
    for i, label in enumerate(rows[0][1]):
        if label in (None, ""):
            continue
        text = str(label).replace("*", "").strip()
        key = by_label.get(_norm(text))
        if key is None:
            for prefix, kind in extra_prefixes:
                if text.startswith(prefix):
                    key = (kind, text[len(prefix):].strip())
                    break
        if key is None:
            report.warn("unknown_column", f"The column '{text}' is not part of the template and is ignored.", sheet=sheet, column=text)
            continue
        if key in seen:
            report.error("duplicate_column", f"The column '{text}' appears twice.", sheet=sheet, column=text)
        seen.add(key)
        mapping[i] = key
    for c in columns:
        if c.required and c.key not in seen:
            report.error("missing_column", f"The required column '{c.label}' is missing.", sheet=sheet, column=c.label)
    return mapping


def _rows_to_dicts(rows, mapping):
    out = []
    for n, values in rows[1:]:
        cells = {}
        for i, key in mapping.items():
            cells[key] = values[i] if i < len(values) else None
        out.append(Row(n, cells))
    return out


# ---------------------------------------------------------------- value parsing


_BAD = object()  # returned for a value that was reported as an error (None means "empty")


def _decimal(report, row, sheet, label, raw, *, places, max_digits=18, allow_zero=True):
    text = _text(raw).replace(",", ".") if not isinstance(raw, (int, float, Decimal)) else str(raw)
    if text == "":
        return None
    try:
        d = Decimal(text)
        if not d.is_finite():
            raise InvalidOperation
    except InvalidOperation:
        report.error("invalid_number", f"'{text}' is not a valid number.", sheet=sheet, row=row, column=label, value=text)
        return _BAD
    if d < 0 or (d == 0 and not allow_zero):
        report.error("negative_number" if d < 0 else "zero_not_allowed",
                     "The value cannot be negative." if d < 0 else "The value must be greater than zero.",
                     sheet=sheet, row=row, column=label, value=text)
        return _BAD
    if -d.as_tuple().exponent > places:
        report.error("too_many_decimals", f"At most {places} decimal places are allowed.", sheet=sheet, row=row, column=label, value=text)
        return _BAD
    if len(d.as_tuple().digits) > max_digits:
        report.error("number_too_big", "The number is too large.", sheet=sheet, row=row, column=label, value=text)
        return _BAD
    return d


def _bool(report, row, sheet, label, raw, default):
    text = _norm(_text(raw))
    if text == "":
        return default
    if text in _TRUE:
        return True
    if text in _FALSE:
        return False
    report.error("invalid_boolean", f"Write '{YES}' or '{NO}'.", sheet=sheet, row=row, column=label, value=_text(raw))
    return default


def _attributes(report, row, sheet, label, raw):
    text = _text(raw)
    out = {}
    for part in filter(None, (p.strip() for p in _ATTR_SPLIT.split(text))):
        key, sep, value = part.partition("=")
        if not sep or not key.strip() or not value.strip():
            report.error("invalid_attributes", "Write attributes as name=value, separated by commas.",
                         sheet=sheet, row=row, column=label, value=part)
            return {}
        out[key.strip()] = value.strip()
    return out


# ---------------------------------------------------------------- validation


@dataclass
class VariantPlan:
    row: int
    sku: str
    data: dict
    prices: list


@dataclass
class ProductPlan:
    row: int
    ref: str
    data: dict
    publish: bool
    images: list
    variants: list = field(default_factory=list)


def _label(key):
    return next((c.label for c in PRODUCT_COLUMNS if c.key == key), key)


def _validate_products(rows, ctx, report):
    plans, by_ref, file_skus, file_barcodes, row_plan = [], {}, {}, {}, {}
    price_cols = {k for r in rows for k in r.cells if isinstance(k, tuple)}
    price_lists = {}
    for kind, name in sorted(price_cols):
        pl = ctx.price_lists.get(_norm(name))
        if pl is None:
            report.error("unknown_price_list", f"There is no price list named '{name}'.", sheet=SHEET_PRODUCTS,
                         column=(PRICE_PREFIX if kind == "price" else COMPARE_PREFIX) + name, value=name)
        elif not pl.is_active:
            report.error("price_list_inactive", f"The price list '{pl.name}' is inactive.", sheet=SHEET_PRODUCTS, column=name, value=name)
        else:
            price_lists[name] = pl

    for r in rows:
        c = r.cells
        sku = _text(c.get("sku"))
        ref = _text(c.get("ref"))
        where = {"sheet": SHEET_PRODUCTS, "row": r.n}
        if not sku:
            report.error("required", "The SKU is required.", column=_label("sku"), **where)
            continue
        if len(sku) > 64:
            report.error("too_long", "At most 64 characters.", column=_label("sku"), value=sku[:40], **where)
        if sku.casefold() in file_skus:
            report.error("duplicate_sku_in_file", f"The SKU '{sku}' is repeated (also in row {file_skus[sku.casefold()]}).",
                         column=_label("sku"), value=sku, **where)
        elif sku.casefold() in ctx.sku_taken:
            report.error("sku_taken", f"The SKU '{sku}' already exists in the system.", column=_label("sku"), value=sku, **where)
        file_skus.setdefault(sku.casefold(), r.n)

        barcode = _text(c.get("barcode"))
        if barcode:
            if len(barcode) > 64:
                report.error("too_long", "At most 64 characters.", column=_label("barcode"), value=barcode[:40], **where)
            if barcode in file_barcodes:
                report.error("duplicate_barcode_in_file", f"The barcode '{barcode}' is repeated (also in row {file_barcodes[barcode]}).",
                             column=_label("barcode"), value=barcode, **where)
            elif barcode in ctx.barcode_taken:
                report.error("barcode_taken", f"The barcode '{barcode}' already exists in the system.",
                             column=_label("barcode"), value=barcode, **where)
            file_barcodes.setdefault(barcode, r.n)

        weight = _decimal(report, r.n, SHEET_PRODUCTS, _label("weight"), c.get("weight"), places=3, max_digits=10)
        reorder = _decimal(report, r.n, SHEET_PRODUCTS, _label("reorder"), c.get("reorder"), places=3, max_digits=14)
        active = _bool(report, r.n, SHEET_PRODUCTS, _label("active"), c.get("active"), True)
        attrs = _attributes(report, r.n, SHEET_PRODUCTS, _label("attributes"), c.get("attributes"))
        vname = _text(c.get("variant_name"))
        if len(vname) > 200:
            report.error("too_long", "At most 200 characters.", column=_label("variant_name"), value=vname[:40], **where)

        prices = []
        for name, pl in price_lists.items():
            col = PRICE_PREFIX + name
            price = _decimal(report, r.n, SHEET_PRODUCTS, col, c.get(("price", name)), places=4)
            compare = _decimal(report, r.n, SHEET_PRODUCTS, COMPARE_PREFIX + name, c.get(("compare", name)), places=4)
            if price is _BAD or compare is _BAD:
                continue
            if price is None:
                if compare is not None:
                    report.error("compare_without_price", "A price is required when a price-before-discount is given.",
                                 column=COMPARE_PREFIX + name, **where)
                continue
            if compare is not None and compare < price:
                report.error("invalid_compare_price", "The price before discount cannot be lower than the price.",
                             column=COMPARE_PREFIX + name, value=_text(c.get(("compare", name))), **where)
                continue
            prices.append({"price_list": pl, "price": price, "compare_at_price": compare})

        variant = VariantPlan(r.n, sku, {
            "sku": sku, "barcode": barcode, "name": vname, "attributes": attrs, "is_active": active,
            **({"weight": weight} if weight is not None and weight is not _BAD else {}),
            **({"reorder_level": reorder} if reorder is not None and reorder is not _BAD else {}),
        }, prices)

        if _text(c.get("images")):
            report.error("image_not_embedded", "The Images column takes the picture itself, not a file name: paste the picture "
                         "over this cell.", column=_label("images"), value=_text(c.get("images"))[:40], **where)

        # product-level cells
        product_cells = {k: _text(c.get(k)) for k in PRODUCT_LEVEL}
        plan = by_ref.get(ref.casefold()) if ref else None
        if plan is not None:
            for k in PRODUCT_LEVEL:
                first = plan.cells[k]
                if product_cells[k] and _norm(product_cells[k]) != _norm(first):
                    report.error("product_field_mismatch",
                                 f"This product-level value differs from the first row of '{ref}' (row {plan.plan.row}). "
                                 "Fill product columns only in the first row.", column=_label(k), value=product_cells[k][:40], **where)
            plan.plan.variants.append(variant)
            row_plan[r.n] = plan.plan
            continue

        # first row of a product: resolve and validate its own fields
        name = product_cells["name"]
        if not name:
            report.error("required", "The product name is required in the first row of a product.", column=_label("name"), **where)
        for k in ("name", "name_ar", "name_en", "name_he"):
            if len(product_cells[k]) > 200:
                report.error("too_long", "At most 200 characters.", column=_label(k), value=product_cells[k][:40], **where)
        brand = category = None
        if product_cells["brand"]:
            found = ctx.find_brand(product_cells["brand"])
            if not found:
                report.error("unknown_brand", f"There is no brand named '{product_cells['brand']}'. Add it first or fix the name.",
                             column=_label("brand"), value=product_cells["brand"], **where)
            elif len(found) > 1:
                report.error("ambiguous_brand", f"More than one brand matches '{product_cells['brand']}'.",
                             column=_label("brand"), value=product_cells["brand"], **where)
            elif not found[0].is_active:
                report.error("brand_inactive", f"The brand '{found[0].name}' is inactive.", column=_label("brand"),
                             value=product_cells["brand"], **where)
            else:
                brand = found[0]
        if product_cells["category"]:
            found = ctx.find_category(product_cells["category"])
            if not found:
                report.error("unknown_category", f"There is no category named '{product_cells['category']}'. "
                             f"Add it first or fix the name (sub-categories: Parent {PATH_SEPARATOR} Child).",
                             column=_label("category"), value=product_cells["category"], **where)
            elif len(found) > 1:
                report.error("ambiguous_category", f"More than one category matches '{product_cells['category']}'. "
                             f"Write the full path: Parent {PATH_SEPARATOR} Child.", column=_label("category"),
                             value=product_cells["category"], **where)
            elif not found[0].is_active:
                report.error("category_inactive", f"The category '{found[0].name}' is inactive.", column=_label("category"),
                             value=product_cells["category"], **where)
            else:
                category = found[0]
        publish = _bool(report, r.n, SHEET_PRODUCTS, _label("publish"), c.get("publish"), False)
        track = _bool(report, r.n, SHEET_PRODUCTS, _label("track"), c.get("track"), True)
        data = {k: product_cells[k] for k in ("name", "name_ar", "name_en", "name_he", "description", "description_ar",
                                              "description_en", "description_he")}
        data.update(category=category, brand=brand, track_inventory=track)
        plan = ProductPlan(r.n, ref, data, publish, [], [variant])
        plans.append(plan)
        row_plan[r.n] = plan
        if ref:
            by_ref[ref.casefold()] = _Group(plan, product_cells)
    for plan in plans:
        if plan.publish and not any(v.data["is_active"] for v in plan.variants):
            report.error("no_active_variant", "A published product needs at least one active variant.",
                         sheet=SHEET_PRODUCTS, row=plan.row, column=_label("publish"))
    return plans, row_plan


@dataclass
class _Group:
    plan: ProductPlan
    cells: dict


def _attach_embedded(pictures, unplaced, images_col, row_plan, ctx, report):
    """Give every picture found in the workbook to the product of the row it sits on; pictures must be in the Images column."""
    where = {"sheet": SHEET_PRODUCTS, "column": _label("images")}
    if unplaced:
        report.error("image_unplaced", f"{unplaced} picture(s) are not placed on a cell. Put each picture on a cell of the "
                     "Images column.", **where)
    for pic in sorted(pictures, key=lambda x: (x.row, x.order)):
        spot = {**where, "row": pic.row}
        if images_col is None or pic.col != images_col:
            report.error("image_wrong_column", "The picture is not in the Images column. Move it into that column.", **spot)
            continue
        plan = row_plan.get(pic.row)
        if plan is None:
            report.error("image_without_row", "The picture is on a row that has no product (it must be on a row with an SKU).", **spot)
            continue
        problem = ctx.check_bytes(pic.data)
        if problem:
            report.error(problem[0], problem[1], **spot)
            continue
        plan.images.append(ImageRef(f"row {pic.row}", pic.data))


def _check_image_counts(plans, report):
    for p in plans:
        if len(p.images) > MAX_IMAGES_PER_PRODUCT:
            report.error("too_many_images", f"At most {MAX_IMAGES_PER_PRODUCT} pictures per product.",
                         sheet=SHEET_PRODUCTS, row=p.row, column=_label("images"))


def _validate_stock(rows, ctx, row_plan, report, can_adjust):
    """The stock columns of the Products sheet (warehouse + quantity + unit cost, per variant row).
    -> {warehouse: [(sku, quantity, unit_cost)]}"""
    out = {}
    keys = ("warehouse", "quantity", "unit_cost")
    given = [r for r in rows if any(_text(r.cells.get(k)) for k in keys)]
    if given and not can_adjust:
        raise ServiceError("forbidden", "Importing stock needs the inventory.adjust permission.", status.HTTP_403_FORBIDDEN)
    for r in given:
        c = r.cells
        where = {"sheet": SHEET_PRODUCTS, "row": r.n}
        plan = row_plan.get(r.n)
        sku, wname = _text(c.get("sku")), _text(c.get("warehouse"))
        if plan is None:  # the row itself was refused (no SKU ...): that error is already reported
            continue
        ok = True
        if not plan.data["track_inventory"]:
            report.error("not_tracked", f"The product of '{sku}' does not track inventory, so it cannot hold stock. "
                         "Set 'Track inventory' to Yes or clear the stock columns.", column=_label("warehouse"), value=sku, **where)
            ok = False
        warehouse = None
        if not wname:
            report.error("required", "The warehouse is required when a quantity is given.", column=_label("warehouse"), **where)
            ok = False
        else:
            found = ctx.find_warehouse(wname)
            if not found:
                report.error("unknown_warehouse", f"There is no warehouse named '{wname}'. Add it first or fix the name.",
                             column=_label("warehouse"), value=wname, **where)
                ok = False
            elif len(found) > 1:
                report.error("ambiguous_warehouse", f"More than one warehouse matches '{wname}'.", column=_label("warehouse"),
                             value=wname, **where)
                ok = False
            elif not found[0].is_active:
                report.error("warehouse_inactive", f"The warehouse '{found[0].name}' is inactive.", column=_label("warehouse"),
                             value=wname, **where)
                ok = False
            else:
                warehouse = found[0]
        qty = _decimal(report, r.n, SHEET_PRODUCTS, _label("quantity"), c.get("quantity"), places=3, max_digits=14, allow_zero=False)
        cost = _decimal(report, r.n, SHEET_PRODUCTS, _label("unit_cost"), c.get("unit_cost"), places=4)
        if qty is None:
            report.error("required", "The quantity is required when a warehouse or cost is given.", column=_label("quantity"), **where)
        if cost is None:
            report.error("required", "The unit cost is required with a quantity.", column=_label("unit_cost"), **where)
        if not ok or qty is None or qty is _BAD or cost is None or cost is _BAD:
            continue
        out.setdefault(warehouse, []).append((sku, qty, cost))
    return out


def _warnings(plans, stock, ctx, report):
    stocked = {sku.casefold() for lines in stock.values() for sku, _q, _c in lines}
    for p in plans:
        where = {"sheet": SHEET_PRODUCTS, "row": p.row, "value": p.data["name"][:40]}
        if not p.images:
            report.warn("no_image", "This product has no picture.", **where)
        if not p.publish:
            continue
        if not any(v.prices for v in p.variants):
            report.warn("published_without_price", "Published but no variant has a price, so customers will not see it.", **where)
        if not any(v.sku.casefold() in stocked for v in p.variants):
            report.warn("published_without_stock", "Published but no stock is imported for it, so customers will not see it "
                        "until stock is added.", **where)


# ---------------------------------------------------------------- entry point


def run(*, actor, workbook, dry_run=True, can_adjust_stock=False, request=None):
    """Validate (and, unless `dry_run`, import) the workbook. Returns the Report; a real run with errors writes nothing."""
    report = Report()
    wb, raw = _open_workbook(workbook)
    try:
        ws = _sheet(wb, SHEET_PRODUCTS, "products")
        if ws is None:
            raise ServiceError("missing_sheet", f"The workbook has no '{SHEET_PRODUCTS}' sheet (use the downloaded template).")
        ctx = Context()
        rows = _read_rows(ws)
        prefixes = ((COMPARE_PREFIX, "compare"), (PRICE_PREFIX, "price"))  # longest first: "السعر قبل الخصم - " also starts with "السعر"
        mapping = _header_map(rows, PRODUCT_COLUMNS, SHEET_PRODUCTS, report, prefixes)
        for title in wb.sheetnames:  # the old separate stock sheet is no longer read: say so instead of silently ignoring it
            if _norm(title) in {_norm("المخزون"), "stock"}:
                report.warn("stock_sheet_ignored", f"The sheet '{title}' is ignored. Put the warehouse, quantity and unit cost "
                            f"in the '{SHEET_PRODUCTS}' sheet.", sheet=title)
        if mapping is None or any(e.code in {"missing_column", "duplicate_column", "missing_header"} for e in report.errors):
            return report
        product_rows = _rows_to_dicts(rows, mapping)
        if not product_rows:
            report.error("empty_file", "The Products sheet has no rows.", sheet=SHEET_PRODUCTS)
            return report
        plans, row_plan = _validate_products(product_rows, ctx, report)
        floating, unplaced = floating_pictures(ws)
        try:
            in_cell = in_cell_pictures(raw, ws.title)
        except Exception:
            in_cell = []
            report.warn("in_cell_pictures_unreadable", "Pictures placed inside cells could not be read; "
                        "paste the pictures over the cells instead.", sheet=SHEET_PRODUCTS)
        images_col = next((i for i, k in mapping.items() if k == "images"), None)
        _attach_embedded(floating + in_cell, unplaced, images_col, row_plan, ctx, report)
        _check_image_counts(plans, report)
        stock = _validate_stock(product_rows, ctx, row_plan, report, can_adjust_stock)
        report.products = len(plans)
        report.variants = sum(len(p.variants) for p in plans)
        report.images = sum(len(p.images) for p in plans)
        report.stock_lines = sum(len(v) for v in stock.values())
        _warnings(plans, stock, ctx, report)
        if report.errors or dry_run:
            return report
        _write(actor, plans, stock, ctx, request, report)
        return report
    finally:
        wb.close()


def _write(actor, plans, stock, ctx, request, report):
    from apps.inventory import services as inventory
    from apps.inventory.models import AdjustmentReason

    saved_files = []
    try:
        with transaction.atomic():
            variants = {}
            for p in plans:
                product = services.create_product(
                    actor=actor, data=dict(p.data),
                    variants=[{**v.data, "prices": v.prices} for v in p.variants], request=request)
                for v in product.variants.all():
                    variants[v.sku.casefold()] = v
                if p.publish:
                    services.publish(actor=actor, product=product, request=request)
                    services.set_visibility(actor=actor, product=product, visible=True, request=request)
                for ref in p.images:
                    image = services.add_image(actor=actor, product=product, file=ContentFile(ref.data, name=ref.label),
                                               alt=product.name, request=request)
                    saved_files.append(image.image.name)
            for warehouse, lines in stock.items():
                adjustment = inventory.create_adjustment(
                    actor=actor, warehouse=warehouse, reason=AdjustmentReason.OPENING, note="Excel product import",
                    lines=[{"variant": variants[sku.casefold()], "counted_qty": qty, "unit_cost": cost} for sku, qty, cost in lines],
                    request=request)
                inventory.post_adjustment(actor=actor, adjustment=adjustment, request=request)
            audit.log("products.imported", actor=actor, request=request, products=len(plans), variants=report.variants,
                      images=report.images, stock_lines=report.stock_lines)
            report.created = len(plans)
    except Exception:
        # the database rolled back; the pictures already written to storage must not be left behind
        from django.core.files.storage import default_storage

        for name in saved_files:
            default_storage.delete(name)
        raise
