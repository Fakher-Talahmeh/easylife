"""The Excel product-import template: the one place that defines every sheet and column.

`import_products.py` reads files with exactly this layout, so the template and the importer cannot drift apart.
Headers are Arabic (the shop's working language); the English key is also accepted when reading.
"""

import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Brand, Category, PriceList

SHEET_PRODUCTS = "المنتجات"
SHEET_HELP = "تعليمات"
SHEET_LISTS = "القوائم"

PRICE_PREFIX = "السعر - "
COMPARE_PREFIX = "السعر قبل الخصم - "
YES, NO = "نعم", "لا"
PATH_SEPARATOR = ">"


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    required: bool
    width: int
    help: str
    example: str = ""


PRODUCT_COLUMNS = [
    Column("ref", "مرجع المنتج", False, 16,
           "أي نص يربط عدّة صفوف (متغيّرات) بمنتج واحد، مثل: SHIRT1. الصفوف التي لها نفس المرجع تصير منتجاً واحداً بعدة متغيّرات. "
           "اتركه فارغاً إذا كان المنتج بمتغيّر واحد.", "SHIRT1"),
    Column("name", "اسم المنتج", True, 28, "الاسم الافتراضي للمنتج (حتى 200 حرف). مطلوب في أول صف من المنتج.", "قميص قطني"),
    Column("name_ar", "الاسم بالعربية", False, 24, "اختياري. إذا تُرك فارغاً يظهر الاسم الافتراضي.", ""),
    Column("name_en", "الاسم بالإنجليزية", False, 24, "اختياري.", "Cotton shirt"),
    Column("name_he", "الاسم بالعبرية", False, 24, "اختياري.", ""),
    Column("description", "الوصف", False, 36, "وصف المنتج الافتراضي. اختياري.", "قميص قطني 100%"),
    Column("description_ar", "الوصف بالعربية", False, 30, "اختياري.", ""),
    Column("description_en", "الوصف بالإنجليزية", False, 30, "اختياري.", ""),
    Column("description_he", "الوصف بالعبرية", False, 30, "اختياري.", ""),
    Column("brand", "الماركة", False, 20,
           "اسم الماركة كما هو في النظام (صفحة الماركات) — اختر من القائمة المنسدلة. لا تُنشأ ماركة جديدة: "
           "إن لم تكن موجودة أضفها أولاً من صفحة الماركات.", "Acme"),
    Column("category", "القسم", False, 24,
           "اسم القسم كما هو في النظام. للأقسام الفرعية استعمل المسار: الأب > الابن (اختر من القائمة المنسدلة). "
           "لا يُنشأ قسم جديد: أضفه أولاً من صفحة الأقسام.", "ملابس > قمصان"),
    Column("publish", "نشر للزبائن", False, 12,
           "نعم = يُنشر ويظهر في المتجر (بشرط وجود سعر ومخزون). لا (أو فارغ) = يُحفظ كمسودة مخفية.", "نعم"),
    Column("track", "تتبع المخزون", False, 12,
           "نعم (الافتراضي) = للمنتج مخزون وكميات. لا = منتج بلا مخزون (لن يظهر للزبائن لأن المنتج الظاهر يلزمه مخزون).", "نعم"),
    Column("images", "الصور", False, 36,
           "ضع صورة المنتج داخل هذا العمود نفسه: انسخ الصورة والصقها فوق خلية هذا العمود في صف المنتج (أو إدراج ← صورة، أو "
           "«وضع في الخلية» في Excel 365). للمنتج بعدة متغيّرات تُحسب الصور الموضوعة في أي صف من صفوفه، وتكون الأولى هي الرئيسية. "
           "حتى 10 صور للمنتج، بصيغة JPG أو PNG أو WebP وحتى 5 MB للصورة. يجب أن تكون الصورة في هذا العمود وفي صف له SKU. "
           "لا تكتب اسم ملف: الصق الصورة نفسها.", "[صورة داخل الخلية]"),
    Column("sku", "رمز الصنف (SKU)", True, 18, "رمز فريد لكل متغيّر (حتى 64 حرفاً). لا يتكرر في الملف ولا مع منتج موجود. مطلوب.", "SH-BL-M"),
    Column("barcode", "الباركود", False, 18, "اختياري. فريد إن وُجد.", "6281000000012"),
    Column("variant_name", "اسم المتغيّر", False, 18, "مثل: أزرق / M. اختياري.", "أزرق / M"),
    Column("attributes", "الخصائص", False, 28, "أزواج اسم=قيمة مفصولة بفاصلة. مثل: اللون=أزرق, المقاس=M", "اللون=أزرق, المقاس=M"),
    Column("weight", "الوزن", False, 10, "رقم (حتى 3 منازل عشرية). اختياري.", "0.25"),
    Column("reorder", "حد إعادة الطلب", False, 14, "الكمية التي يُنبَّه عندها لإعادة الطلب. اختياري (الافتراضي 0).", "5"),
    Column("active", "فعّال", False, 10, "نعم (الافتراضي) أو لا.", "نعم"),
    Column("warehouse", "المخزن", False, 22,
           "اسم المخزن الذي تدخل إليه الكمية، كما هو في النظام (اختر من القائمة المنسدلة). لا يُنشأ مخزن جديد. "
           "الأعمدة الثلاثة (المخزن والكمية وتكلفة الوحدة) تُملأ معاً أو تُترك فارغة كلها. تتطلب صلاحية تعديل المخزون.", "Main warehouse"),
    Column("quantity", "الكمية", False, 12, "الكمية الافتتاحية في هذا المخزن (أكبر من صفر، حتى 3 منازل عشرية). المنتج بلا كمية لا يظهر للزبائن.", "20"),
    Column("unit_cost", "تكلفة الوحدة", False, 14,
           "تكلفة شراء الوحدة بالعملة الأساسية (للمحاسبة). مطلوبة مع الكمية.", "12.5"),
]

_HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
_REQ_FILL = PatternFill("solid", fgColor="C00000")
_HEAD_FONT = Font(bold=True, color="FFFFFF")


def category_paths():
    """{category: "Parent > Child"} for every category (the form people pick from)."""
    cats = {c.pk: c for c in Category.objects.all()}

    def path(c):
        parts = []
        while c is not None:
            parts.append(c.name)
            c = cats.get(c.parent_id)
        return f" {PATH_SEPARATOR} ".join(reversed(parts))

    return {c: path(c) for c in sorted(cats.values(), key=lambda c: (c.sort_order, c.name))}


def price_columns(price_lists):
    """Two extra columns per active price list: its price and its compare-at ("before discount") price."""
    out = []
    for pl in price_lists:
        out.append((PRICE_PREFIX + pl.name, f"سعر البيع في قائمة «{pl.name}» ({pl.currency.code}). اتركه فارغاً إن لم يُسعَّر الصنف فيها. "
                    "الصنف الذي بلا سعر لا يظهر للزبائن.", 16))
        out.append((COMPARE_PREFIX + pl.name, f"اختياري: السعر قبل الخصم في قائمة «{pl.name}»، ويجب ألا يقل عن السعر.", 16))
    return out


def _header(ws, labels, widths, required_flags):
    for i, (label, width, required) in enumerate(zip(labels, widths, required_flags), start=1):
        cell = ws.cell(row=1, column=i, value=label + (" *" if required else ""))
        cell.font = _HEAD_FONT
        cell.fill = _REQ_FILL if required else _HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    ws.sheet_view.rightToLeft = True


def _dropdown(ws, col_index, values_range, *, rows=2000, strict=False):
    dv = DataValidation(type="list", formula1=values_range, allow_blank=True, showDropDown=False)
    dv.errorStyle = "stop" if strict else "warning"
    dv.errorTitle, dv.error = "قيمة غير موجودة", "هذه القيمة ليست في القائمة. تأكد من الاسم كما هو في النظام."
    dv.showErrorMessage = True
    letter = get_column_letter(col_index)
    dv.add(f"{letter}2:{letter}{rows + 1}")
    ws.add_data_validation(dv)


def build_template():
    """The blank workbook people fill in: ONE data sheet (Products) plus read-only help: instructions and the live lists."""
    from apps.inventory.models import Warehouse

    brands = list(Brand.objects.filter(is_active=True).order_by("name"))
    cats = category_paths()
    warehouses = list(Warehouse.objects.filter(is_active=True).order_by("code"))
    price_lists = list(PriceList.objects.filter(is_active=True).select_related("currency").order_by("-is_default", "name"))
    extra = price_columns(price_lists)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_PRODUCTS
    labels = [c.label for c in PRODUCT_COLUMNS] + [e[0] for e in extra]
    widths = [c.width for c in PRODUCT_COLUMNS] + [e[2] for e in extra]
    flags = [c.required for c in PRODUCT_COLUMNS] + [False] * len(extra)
    _header(ws, labels, widths, flags)

    lists = wb.create_sheet(SHEET_LISTS)
    lists.sheet_view.rightToLeft = True
    headers = ["الماركات", "الأقسام", "المخازن", "قوائم الأسعار"]
    for i, h in enumerate(headers, start=1):
        c = lists.cell(row=1, column=i, value=h)
        c.font, c.fill = _HEAD_FONT, _HEAD_FILL
        lists.column_dimensions[get_column_letter(i)].width = 34
    for r, b in enumerate(brands, start=2):
        lists.cell(row=r, column=1, value=b.name)
    for r, path in enumerate(cats.values(), start=2):
        lists.cell(row=r, column=2, value=path)
    for r, w in enumerate(warehouses, start=2):
        lists.cell(row=r, column=3, value=w.name)
    for r, pl in enumerate(price_lists, start=2):
        lists.cell(row=r, column=4, value=f"{pl.name} ({pl.currency.code}){' — الافتراضية' if pl.is_default else ''}")

    keys = [c.key for c in PRODUCT_COLUMNS]
    q = f"'{SHEET_LISTS}'!"
    if brands:
        _dropdown(ws, keys.index("brand") + 1, f"={q}$A$2:$A${len(brands) + 1}")
    if cats:
        _dropdown(ws, keys.index("category") + 1, f"={q}$B$2:$B${len(cats) + 1}")
    if warehouses:
        _dropdown(ws, keys.index("warehouse") + 1, f"={q}$C$2:$C${len(warehouses) + 1}")
    for key in ("publish", "track", "active"):
        _dropdown(ws, keys.index(key) + 1, f'"{YES},{NO}"', strict=True)

    _write_help(wb, extra)
    wb.move_sheet(SHEET_HELP, offset=-(len(wb.sheetnames) - 1))  # instructions first, so it is read first
    wb.active = wb.sheetnames.index(SHEET_PRODUCTS)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _write_help(wb, extra):
    ws = wb.create_sheet(SHEET_HELP)
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90
    ws.column_dimensions["D"].width = 30
    steps = [
        "كيف تستورد المنتجات",
        "1) املأ ورقة «المنتجات»: كل صف = متغيّر واحد (صنف) في مخزن واحد. المنتج بعدة متغيّرات: أعطها نفس «مرجع المنتج».",
        "2) كل البيانات في ورقة «المنتجات» وحدها (بما فيها المخزن والكمية والتكلفة). ورقتا «تعليمات» و«القوائم» للمساعدة فقط ولا يُقرآن عند الاستيراد.",
        "3) الأعمدة بعلامة * مطلوبة. الماركة والقسم والمخزن تُكتب بأسمائها كما في النظام (توجد قوائم منسدلة وورقة «القوائم»).",
        "4) الصور: الصق صورة كل منتج داخل خلية عمود «الصور» في صفه (الصور تُحفظ داخل ملف Excel نفسه، ولا حاجة لرفع شيء آخر).",
        "5) في صفحة الاستيراد اضغط «فحص الملف» أولاً: يُعرض لك كل خطأ مع رقم الصف واسم العمود، ولا يُحفظ شيء. ثم اضغط «استيراد».",
        "6) الاستيراد كله أو لا شيء: إن وُجد أي خطأ لا يُضاف أي منتج. الاستيراد يضيف منتجات جديدة فقط ولا يعدّل الموجود (الـ SKU الموجود يُرفض).",
        "",
        "الحدود: حتى 2000 صف، وملف Excel حتى 60 MB (بما فيه الصور)، والصورة الواحدة حتى 5 MB، وحتى 10 صور للمنتج.",
        "",
    ]
    for r, text in enumerate(steps, start=1):
        c = ws.cell(row=r, column=1, value=text)
        c.alignment = Alignment(wrap_text=False, vertical="center")
        if r == 1:
            c.font = Font(bold=True, size=14)
    row = len(steps) + 1
    for i, h in enumerate(["الورقة / العمود", "مطلوب", "الشرح والقواعد", "مثال"], start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill = _HEAD_FONT, _HEAD_FILL

    def put(sheet, col):
        nonlocal row
        row += 1
        ws.cell(row=row, column=1, value=f"{sheet} ← {col.label}")
        ws.cell(row=row, column=2, value="نعم" if col.required else "لا")
        c = ws.cell(row=row, column=3, value=col.help)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=4, value=col.example)

    for col in PRODUCT_COLUMNS:
        put(SHEET_PRODUCTS, col)
    for label, helptext, _w in extra:
        row += 1
        ws.cell(row=row, column=1, value=f"{SHEET_PRODUCTS} ← {label}")
        ws.cell(row=row, column=2, value="لا")
        c = ws.cell(row=row, column=3, value=helptext)
        c.alignment = Alignment(wrap_text=True, vertical="top")
