import io
import re
import zipfile

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.utils import get_column_letter
from PIL import Image
from rest_framework.test import APIClient

from apps.catalog.import_template import COMPARE_PREFIX, PRICE_PREFIX, PRODUCT_COLUMNS, SHEET_PRODUCTS
from apps.catalog.models import Brand, Category, Product, ProductImage, ProductVariant
from apps.catalog.tests.helpers import price_list, warehouse
from apps.core.testing import bootstrap, client_for, make_admin, make_customer, make_staff
from apps.inventory.models import StockLevel, StockMovement

B = "/api/v1/admin/products/"
LABEL = {c.key: c.label for c in PRODUCT_COLUMNS}


@pytest.fixture(autouse=True)
def _setup(db, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    bootstrap()


@pytest.fixture
def admin():
    return client_for(make_admin())


@pytest.fixture
def shop(db):
    """The related records an import refers to by name."""
    price_list("ILS")
    brand = Brand.objects.create(name="Acme", name_ar="أكمي")
    root = Category.objects.create(name="Clothes", slug="clothes")
    shirts = Category.objects.create(name="Shirts", slug="shirts", parent=root)
    toys = Category.objects.create(name="Toys", slug="toys")
    return {"brand": brand, "root": root, "shirts": shirts, "toys": toys, "wh": warehouse("MAIN", is_default=True)}


def _header(key):
    if key.startswith("price:"):
        return PRICE_PREFIX + key[6:]
    if key.startswith("compare:"):
        return COMPARE_PREFIX + key[8:]
    return LABEL[key]


def xlsx(products, *, name="products.xlsx", keep_header=True, pictures=(), in_cell=()):
    """products: list of {key: value}. Keys: the column keys, 'price:<list name>', 'compare:<list name>'.
    pictures: [(row, column key, png bytes)] floating pictures anchored on that cell (what Insert > Picture makes);
    in_cell: [(row, column key, bytes)] pictures placed inside the cell (Excel 365 "Place in Cell")."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_PRODUCTS
    keys = list(dict.fromkeys(k for row in products for k in row))
    if keep_header:
        ws.append([_header(k) for k in keys])
    for row in products:
        ws.append([row.get(k) for k in keys])
    for row, key, data in pictures:
        ws.add_image(XLImage(io.BytesIO(data)), f"{get_column_letter(keys.index(key) + 1)}{row}")
    for row, key, _data in in_cell:
        ws[f"{get_column_letter(keys.index(key) + 1)}{row}"] = "placeholder"
    out = io.BytesIO()
    wb.save(out)
    content = out.getvalue()
    if in_cell:
        content = _place_in_cells(content, [(f"{get_column_letter(keys.index(k) + 1)}{r}", d) for r, k, d in in_cell])
    return SimpleUploadedFile(name, content, content_type="application/vnd.ms-excel")


def _place_in_cells(content, cells):
    """Rewrite a workbook the way Excel 365 stores a picture placed in a cell: a rich value per picture (media + rel +
    value structure + metadata) and `vm`/`t="e"` on the cell."""
    src = zipfile.ZipFile(io.BytesIO(content))
    files = {n: src.read(n) for n in src.namelist()}
    sheet = files["xl/worksheets/sheet1.xml"].decode()
    rvs, rels, relx = [], [], []
    for i, (ref, data) in enumerate(cells):
        files[f"xl/media/inimage{i + 1}.png"] = data
        rels.append(f'<Relationship Id="rId{i + 1}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
                    f'Target="../media/inimage{i + 1}.png"/>')
        relx.append(f'<rel xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId{i + 1}"/>')
        rvs.append(f'<rv s="0"><v>{i}</v><v>5</v></rv>')
        cell = re.compile(rf'<c r="{ref}"[^>]*?(/>|>.*?</c>)', re.S)
        assert cell.search(sheet), ref
        sheet = cell.sub(f'<c r="{ref}" vm="{i + 1}" t="e"><v>#VALUE!</v></c>', sheet, count=1)
    n = len(cells)
    files["xl/worksheets/sheet1.xml"] = sheet.encode()
    files["xl/richData/rdrichvalue.xml"] = (
        f'<rvData xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="{n}">{"".join(rvs)}</rvData>').encode()
    files["xl/richData/rdrichvaluestructure.xml"] = (
        '<rvStructures xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="1">'
        '<s t="_localImage"><k n="_rvRel:LocalImageIdentifier" t="i"/><k n="CalcOrigin" t="i"/></s></rvStructures>').encode()
    files["xl/richData/richValueRel.xml"] = (
        f'<richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel">{"".join(relx)}</richValueRels>').encode()
    files["xl/richData/_rels/richValueRel.xml.rels"] = (
        f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{"".join(rels)}</Relationships>').encode()
    bks = "".join(f'<bk><extLst><ext uri="{{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}}"><xlrd:rvb i="{i}"/></ext></extLst></bk>' for i in range(n))
    vms = "".join(f'<bk><rc t="1" v="{i}"/></bk>' for i in range(n))
    files["xl/metadata.xml"] = (
        '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
        '<metadataTypes count="1"><metadataType name="XLRICHVALUE" minSupportedVersion="120000"/></metadataTypes>'
        f'<futureMetadata name="XLRICHVALUE" count="{n}">{bks}</futureMetadata>'
        f'<valueMetadata count="{n}">{vms}</valueMetadata></metadata>').encode()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name, data in files.items():
            z.writestr(name, data)
    return out.getvalue()


def png(color="red", size=(40, 30)):
    out = io.BytesIO()
    Image.new("RGB", size, color).save(out, "PNG")
    return out.getvalue()


def post(client, file, dry_run=False):
    data = {"file": file, "dry_run": "true" if dry_run else "false"}
    return client.post(B + "import/", data, format="multipart")


def codes(resp):
    return {(e["sheet"], e["row"], e["column"], e["code"]) for e in resp.data["errors"]}


def error_codes(resp):
    return {e["code"] for e in resp.data["errors"]}


GOOD = [
    {"ref": "SHIRT1", "name": "Cotton shirt", "name_en": "Cotton shirt", "brand": "Acme", "category": "Clothes > Shirts",
     "publish": "نعم", "images": None, "sku": "SH-M", "variant_name": "M", "attributes": "size=M, color=blue",
     "price:Retail ILS": 50, "compare:Retail ILS": 70, "warehouse": "Warehouse MAIN", "quantity": 10, "unit_cost": 20},
    {"ref": "SHIRT1", "images": None, "sku": "SH-L", "variant_name": "L", "attributes": "size=L, color=blue", "price:Retail ILS": 55,
     "warehouse": "Warehouse MAIN", "quantity": 4, "unit_cost": 21.5},
    {"name": "Robot", "brand": "أكمي", "category": "Toys", "sku": "ROB-1", "barcode": 6281000000012, "price:Retail ILS": "30.5"},
]
# ------------------------------------------------------------ the template


def test_template_has_every_sheet_column_and_the_live_lists(admin, shop):
    r = admin.get(B + "import/template/")
    assert r.status_code == 200 and "spreadsheetml" in r["Content-Type"] and "attachment" in r["Content-Disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert set(wb.sheetnames) == {"تعليمات", "المنتجات", "القوائم"}  # one data sheet; the others are help only
    header = [c.value for c in wb["المنتجات"][1]]
    assert "اسم المنتج *" in header and "رمز الصنف (SKU) *" in header and "الماركة" in header and "القسم" in header
    assert "السعر - Retail ILS" in header and "السعر قبل الخصم - Retail ILS" in header
    assert {"المخزن", "الكمية", "تكلفة الوحدة", "الصور"} <= set(header)  # stock and pictures are columns of the same sheet
    lists = {c.value: [x.value for x in col[1:] if x.value] for col in wb["القوائم"].iter_cols() for c in col[:1]}
    assert lists["الماركات"] == ["Acme"] and "Clothes > Shirts" in lists["الأقسام"] and lists["المخازن"] == ["Warehouse MAIN"]
    assert any("Retail ILS" in v for v in lists["قوائم الأسعار"])


def test_import_endpoints_need_the_import_permission(shop):
    for who in (client_for(make_staff(["products.view", "products.create", "products.update"])), client_for(make_customer())):
        assert who.get(B + "import/template/").status_code == 403
        assert post(who, xlsx(GOOD)).status_code == 403
    assert APIClient().get(B + "import/template/").status_code == 401
    assert APIClient().post(B + "import/", {}, format="multipart").status_code == 401
    allowed = client_for(make_staff(["products.import"]))
    assert allowed.get(B + "import/template/").status_code == 200


# ------------------------------------------------------------ importing


def test_full_import_creates_products_variants_prices_images_and_stock(admin, shop):
    r = post(admin, xlsx(GOOD, pictures=[(2, "images", png("red")), (3, "images", png("blue"))]))
    assert r.status_code == 201, r.data
    assert r.data["valid"] and r.data["summary"] == {"products": 2, "variants": 3, "images": 2, "stock_lines": 2}

    shirt = Product.objects.get(name="Cotton shirt")
    assert (shirt.brand, shirt.category) == (shop["brand"], shop["shirts"])  # matched by name / path
    assert shirt.status == "active" and shirt.is_visible_on_store and shirt.track_inventory and shirt.name_en == "Cotton shirt"
    m, l = shirt.variants.get(sku="SH-M"), shirt.variants.get(sku="SH-L")
    assert m.attributes == {"size": "M", "color": "blue"} and m.name == "M"
    p = m.prices.get()
    assert (str(p.price), str(p.compare_at_price), p.price_list.name) == ("50.0000", "70.0000", "Retail ILS")
    assert l.prices.get().compare_at_price is None
    imgs = list(shirt.images.all())
    assert len(imgs) == 2 and imgs[0].is_primary and not imgs[1].is_primary and imgs[0].image.name.endswith(".webp")

    robot = Product.objects.get(name="Robot")
    assert robot.brand == shop["brand"] and robot.category == shop["toys"]  # brand by its Arabic name
    assert robot.status == "draft" and not robot.is_visible_on_store and robot.images.count() == 0
    v = robot.variants.get()
    assert v.barcode == "6281000000012" and str(v.prices.get().price) == "30.5000"

    assert StockLevel.objects.get(variant=m, warehouse=shop["wh"]).on_hand == 10
    assert StockLevel.objects.get(variant=l, warehouse=shop["wh"]).on_hand == 4
    m.refresh_from_db()
    assert str(m.cost_price) == "20.0000" and StockMovement.objects.filter(variant=m, type="opening").count() == 1
    # published + priced + stocked -> customers can see it; the draft and the unstocked one cannot
    names = [x["name"] for x in APIClient().get("/api/v1/store/products/").data["results"]]
    assert names == ["Cotton shirt"]
    assert any(w["code"] == "no_image" for w in r.data["warnings"])


def test_dry_run_validates_but_writes_nothing(admin, shop):
    r = post(admin, xlsx(GOOD, pictures=[(2, "images", png("red")), (3, "images", png("blue"))]), dry_run=True)
    assert r.status_code == 200 and r.data["valid"] and r.data["dry_run"] and r.data["summary"]["products"] == 2
    assert Product.objects.count() == 0 and ProductImage.objects.count() == 0 and StockLevel.objects.count() == 0


def test_products_without_stock_columns_import_fine_without_inventory_adjust(shop):
    staff = client_for(make_staff(["products.import"]))
    r = post(staff, xlsx(GOOD[2:]))
    assert r.status_code == 201 and Product.objects.count() == 1
    assert post(staff, xlsx(GOOD[:1])).status_code == 403  # stock needs inventory.adjust


# ------------------------------------------------------------ every mistake is reported with sheet / row / column


def test_all_problems_are_reported_together_and_nothing_is_imported(admin, shop):
    existing = Product.objects.create(name="Old", slug="old")
    ProductVariant.objects.create(product=existing, sku="TAKEN-1", barcode="999")
    bad = [
        {"name": "A", "brand": "Nope", "category": "Clothes > Hats", "sku": "TAKEN-1", "price:Retail ILS": 5},
        {"name": "B", "sku": "DUP", "price:Retail ILS": -3, "publish": "ربما"},
        {"name": "C", "sku": "dup", "price:Retail ILS": 10, "compare:Retail ILS": 4, "barcode": "999"},
        {"name": "D", "sku": "D1", "price:Retail ILS": "abc", "attributes": "colorblue", "images": "missing.png"},
        {"ref": "X", "name": "E", "sku": "E1"},
        {"ref": "X", "name": "Other name", "sku": "E2"},
        {"name": "F", "sku": "F1", "weight": "1.23456"},
        {"sku": "G1"},
        {"name": "H", "sku": "H1", "warehouse": "Ghost", "quantity": 0, "unit_cost": None},
    ]
    r = post(admin, xlsx(bad))
    assert r.status_code == 400 and r.data["code"] == "import_invalid" and not r.data["valid"]
    got = codes(r)
    assert (SHEET_PRODUCTS, 2, "الماركة", "unknown_brand") in got
    assert (SHEET_PRODUCTS, 2, "القسم", "unknown_category") in got
    assert (SHEET_PRODUCTS, 2, "رمز الصنف (SKU)", "sku_taken") in got
    assert (SHEET_PRODUCTS, 3, "السعر - Retail ILS", "negative_number") in got
    assert (SHEET_PRODUCTS, 3, "نشر للزبائن", "invalid_boolean") in got
    assert (SHEET_PRODUCTS, 4, "رمز الصنف (SKU)", "duplicate_sku_in_file") in got
    assert (SHEET_PRODUCTS, 4, "السعر قبل الخصم - Retail ILS", "invalid_compare_price") in got
    assert (SHEET_PRODUCTS, 4, "الباركود", "barcode_taken") in got
    assert (SHEET_PRODUCTS, 5, "السعر - Retail ILS", "invalid_number") in got
    assert (SHEET_PRODUCTS, 5, "الخصائص", "invalid_attributes") in got
    assert (SHEET_PRODUCTS, 5, "الصور", "image_not_embedded") in got
    assert (SHEET_PRODUCTS, 7, "اسم المنتج", "product_field_mismatch") in got
    assert (SHEET_PRODUCTS, 8, "الوزن", "too_many_decimals") in got
    assert (SHEET_PRODUCTS, 9, "اسم المنتج", "required") in got
    assert (SHEET_PRODUCTS, 10, "المخزن", "unknown_warehouse") in got
    assert (SHEET_PRODUCTS, 10, "الكمية", "zero_not_allowed") in got
    assert (SHEET_PRODUCTS, 10, "تكلفة الوحدة", "required") in got
    assert Product.objects.count() == 1  # only the pre-existing one


def test_a_file_name_in_the_images_column_is_refused_because_pictures_must_be_pasted(admin, shop):
    r = post(admin, xlsx([{"name": "A", "sku": "A1", "images": "a.png"}]), dry_run=True)
    assert {(e["column"], e["code"]) for e in r.data["errors"]} == {("الصور", "image_not_embedded")}


def test_bad_files_are_refused_clearly(admin, shop):
    assert post(admin, SimpleUploadedFile("data.csv", b"a,b")).data["code"] == "invalid_file"
    assert post(admin, SimpleUploadedFile("data.xlsx", b"not a workbook")).data["code"] == "invalid_file"
    wb = Workbook()
    wb.active.title = "Other"
    out = io.BytesIO()
    wb.save(out)
    assert post(admin, SimpleUploadedFile("x.xlsx", out.getvalue())).data["code"] == "missing_sheet"
    assert "missing_column" in error_codes(post(admin, xlsx([{"name": "A"}]), dry_run=True))  # no SKU column at all
    unknown_list = [{"name": "A", "sku": "A1", "price:Ghost list": 5}]
    assert "unknown_price_list" in error_codes(post(admin, xlsx(unknown_list), dry_run=True))
    r = post(admin, xlsx([]), dry_run=True)  # a validation run reports problems with 200; only a real import is a 400
    assert r.status_code == 200 and not r.data["valid"] and error_codes(r) == {"missing_header"}
    assert post(admin, xlsx([]), dry_run=False).status_code == 400


def test_a_failure_while_writing_leaves_nothing_behind(admin, shop, monkeypatch):
    from apps.inventory import services as inventory

    def boom(**kw):
        raise RuntimeError("disk on fire")

    monkeypatch.setattr(inventory, "post_adjustment", boom)
    client = APIClient()
    client.force_authenticate(make_admin())
    client.raise_request_exception = False
    r = post(client, xlsx(GOOD, pictures=[(2, "images", png("red")), (3, "images", png("blue"))]))
    assert r.status_code == 500
    assert Product.objects.count() == 0 and ProductImage.objects.count() == 0 and StockLevel.objects.count() == 0
    import os
    from django.conf import settings

    media = os.path.join(settings.MEDIA_ROOT, "products")
    left = [f for _d, _s, fs in os.walk(media) for f in fs] if os.path.isdir(media) else []
    assert left == []  # the pictures written before the failure were deleted


def test_untracked_products_cannot_receive_stock_and_inactive_lists_are_refused(admin, shop):
    rows = [{"name": "Service", "track": "لا", "sku": "SV-1", "warehouse": "MAIN", "quantity": 1, "unit_cost": 1}]
    r = post(admin, xlsx(rows), dry_run=True)
    assert "not_tracked" in error_codes(r)
    price_list("USD", default=False, name="Old list")
    from apps.catalog.models import PriceList

    PriceList.objects.filter(name="Old list").update(is_active=False)
    r = post(admin, xlsx([{"name": "A", "sku": "A1", "price:Old list": 5}]), dry_run=True)
    assert "price_list_inactive" in error_codes(r)
    inactive_wh = warehouse("OLD", is_active=False)
    r = post(admin, xlsx([{"name": "A", "sku": "A1", "warehouse": inactive_wh.name, "quantity": 1, "unit_cost": 1}]), dry_run=True)
    assert "warehouse_inactive" in error_codes(r)


def test_published_products_that_customers_would_not_see_get_a_warning(admin, shop):
    rows = [{"name": "A", "sku": "A1", "publish": "نعم", "images": ""}]
    r = post(admin, xlsx(rows), dry_run=True)
    assert {"published_without_price", "published_without_stock"} <= {w["code"] for w in r.data["warnings"]}
    assert r.data["valid"]


# ------------------------------------------------------------ pictures inside the Excel file (the Images column)


def test_pictures_placed_in_the_images_column_become_the_product_pictures(admin, shop):
    rows = [
        {"ref": "P", "name": "Shirt", "sku": "S-1", "images": None},
        {"ref": "P", "sku": "S-2", "images": None},
        {"name": "Robot", "sku": "R-1", "images": None},
        {"name": "Plain", "sku": "N-1", "images": None},
    ]
    pics = [(2, "images", png("red")), (3, "images", png("blue")), (4, "images", png("green"))]
    r = post(admin, xlsx(rows, pictures=pics))
    assert r.status_code == 201, r.data
    assert r.data["summary"]["images"] == 3
    shirt = Product.objects.get(name="Shirt")
    imgs = list(shirt.images.all())  # a picture on any row of the product belongs to the product, in row order
    assert len(imgs) == 2 and imgs[0].is_primary and not imgs[1].is_primary
    assert Product.objects.get(name="Robot").images.count() == 1 and Product.objects.get(name="Plain").images.count() == 0
    assert [w["value"] for w in r.data["warnings"] if w["code"] == "no_image"] == ["Plain"]


def test_pictures_placed_inside_a_cell_are_read_too(admin, shop):
    rows = [{"name": "Shirt", "sku": "S-1", "images": None}, {"name": "Robot", "sku": "R-1", "images": None}]
    r = post(admin, xlsx(rows, in_cell=[(2, "images", png("red")), (3, "images", png("blue"))]))
    assert r.status_code == 201, r.data
    assert r.data["summary"]["images"] == 2
    assert Product.objects.get(name="Shirt").images.count() == 1 and Product.objects.get(name="Robot").images.count() == 1
    bad = post(admin, xlsx([{"name": "X", "sku": "X-1", "images": None}], in_cell=[(2, "images", b"not a picture")]), dry_run=True)
    assert "invalid_image" in error_codes(bad)


def test_misplaced_pictures_are_reported_not_silently_dropped(admin, shop):
    rows = [{"name": "Shirt", "sku": "S-1", "images": None}]
    good = png("red")
    r = post(admin, xlsx(rows, pictures=[(2, "name", good), (7, "images", good), (1, "images", good)]), dry_run=True)
    got = {(e["row"], e["code"]) for e in r.data["errors"]}
    assert got == {(2, "image_wrong_column"), (7, "image_without_row"), (1, "image_without_row")}
    assert all(e["column"] == "الصور" for e in r.data["errors"])
    assert Product.objects.count() == 0

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_PRODUCTS
    ws.append([LABEL["name"], LABEL["sku"], LABEL["images"]])
    ws.append(["A", "A-1", None])
    img = XLImage(io.BytesIO(good))
    img.anchor = AbsoluteAnchor(pos=XDRPoint2D(100, 100), ext=XDRPositiveSize2D(100, 100))
    ws.add_image(img)
    out = io.BytesIO()
    wb.save(out)
    r = post(admin, SimpleUploadedFile("abs.xlsx", out.getvalue()), dry_run=True)
    assert "image_unplaced" in error_codes(r)


def test_too_many_embedded_pictures_for_one_product(admin, shop):
    rows = [{"ref": "P", "name": "Shirt", "sku": f"S-{i}", "images": None} if i == 0 else {"ref": "P", "sku": f"S-{i}", "images": None}
            for i in range(11)]
    pics = [(i + 2, "images", png()) for i in range(11)]
    assert "too_many_images" in error_codes(post(admin, xlsx(rows, pictures=pics), dry_run=True))


def test_stock_columns_are_all_or_nothing_and_the_old_stock_sheet_is_flagged(admin, shop):
    only_warehouse = [{"name": "A", "sku": "A1", "warehouse": "MAIN"}]
    r = post(admin, xlsx(only_warehouse), dry_run=True)
    assert {(e["column"], e["code"]) for e in r.data["errors"]} == {("الكمية", "required"), ("تكلفة الوحدة", "required")}
    only_qty = [{"name": "A", "sku": "A1", "quantity": 5}]
    assert {e["column"] for e in post(admin, xlsx(only_qty), dry_run=True).data["errors"]} == {"المخزن", "تكلفة الوحدة"}

    wb = load_workbook(io.BytesIO(xlsx([{"name": "A", "sku": "A1"}]).read()))
    wb.create_sheet("المخزون")
    out = io.BytesIO()
    wb.save(out)
    r = post(admin, SimpleUploadedFile("old.xlsx", out.getvalue()), dry_run=True)
    assert r.data["valid"] and "stock_sheet_ignored" in {w["code"] for w in r.data["warnings"]}
